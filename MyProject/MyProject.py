import openpyxl
import datetime
from openpyxl import Workbook
from openpyxl.styles import (PatternFill, Border, Side,
                            Alignment, Font, GradientFill, NamedStyle)
from datetime import date
import locale
locale.setlocale(locale.LC_ALL, 'ru')
mydate = date.today()

wb = Workbook()
wb.create_sheet("Journal")
wb.active = wb["Journal"]
ws = wb.active
print(ws.title)
#Отсутствие сетки
ws.sheet_view.showGridLines = False

#Style "main"
font_size = 14
name_style = NamedStyle(name="main")
name_style.font = Font(size=font_size, bold=True, name='Times New Roman')
bd = Side(style='thick', color="000000")
name_style.border = Border(left=bd, top=bd, right=bd, bottom=bd)
wb.add_named_style(name_style)

#Style "other"
font_size = 14
name_style = NamedStyle(name="other")
name_style.font = Font(size=font_size, bold=False, name='Times New Roman')
bd = Side(style='thin', color="000000")
name_style.border = Border(left=bd, top=bd, right=bd, bottom=bd)
wb.add_named_style(name_style)

# объединим ячейки в диапазоне `A1:A2`
ws.merge_cells('A1:A2')
# в данном случае крайняя верхняя-левая ячейка это `A1`
megre_cell = ws['A1']
# запишем в нее текст
megre_cell.value = 'Номер извещения'
ws["A1"].style = "main"
ws["A2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('B1:B2')
megre_cell = ws['B1']
megre_cell.value = 'Дата'
ws['B1'].style = "main"
ws["B2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('C1:C2')
megre_cell = ws['C1']
megre_cell.value = 'Тип изделия'
ws['C1'].style = "main"
ws["C2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('D1:D2')
megre_cell = ws['D1']
megre_cell.value = 'Номер зав. заказа'
ws['D1'].style = "main"
ws["D2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('E1:E2')
megre_cell = ws['E1']
megre_cell.value = 'Содержание'
ws['E1'].style = "main"
ws["E2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('F1:F2')
megre_cell = ws['F1']
megre_cell.value = 'Ф.И.О. исполнителя'
ws['F1'].style = "main"
ws["F2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('G1:G2')
megre_cell = ws['G1']
megre_cell.value = 'Дата сдачи в архив'
ws['G1'].style = "main"
ws["G2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('K1:K2')
megre_cell = ws['K1']
megre_cell.value = 'Примечание'
ws['K1'].style = "main"
ws["K2"].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws.merge_cells('H1:J1')
megre_cell = ws['H1']
megre_cell.value = 'Согласование'
ws['H1'].style = "main"
megre_cell.fill = PatternFill('solid', fgColor="DAA520")

ws['H2'].value = "C 1 отд."
ws['H2'].style = "main"
ws['H2'].fill = PatternFill('solid', fgColor="DAA520")

ws['I2'].value = "Со 2 отд."
ws['I2'].style = "main"
ws['I2'].fill = PatternFill('solid', fgColor="DAA520")

ws['J2'].value = "С 3 отд."
ws['J2'].style = "main"
ws['J2'].fill = PatternFill('solid', fgColor="DAA520")

ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 17
ws.column_dimensions['D'].width = 25
ws.column_dimensions['E'].width = 30
ws.column_dimensions['F'].width = 25
ws.column_dimensions['G'].width = 25
ws.column_dimensions['H'].width = 12
ws.column_dimensions['I'].width = 12
ws.column_dimensions['J'].width = 12
ws.column_dimensions['K'].width = 20
ws.row_dimensions[1].height = 30

ws.append({1:"100-1", 2:"01.01.2024", 3: "АКГ-1", 4: "10800-0000", 5: "Закупка комплектующих", 6: "Иванов", 7: mydate.strftime("%d.%m.%Y"), 8: "+", 9: "+", 10: "+", 11: "-"})
ws.append({1:"100-2", 2:"01.02.2024", 3: "АКГ-2", 4: "13330-6789", 5: "Изготовление панели", 6: "Петров", 7: mydate.strftime("%d.%m.%Y"), 8: "+", 9: "+", 10: "+", 11: "-"})
ws.append({1:"100-3", 2:"01.03.2024", 3: "АЕМ-1", 4: "12345-3456", 5: "Доработка комплекта", 6: "Сидоров", 7: "02.03.2024", 8: "+", 9: "+", 10: "+", 11: "-"})
ws.append({1:"100-4", 2:"01.04.2024", 3: "АКГ-1", 4: "10800-0000", 5: "Отгрузка комплекта", 6: "Семенов", 7: "25.06.2024", 8: "+", 9: "+", 10: "+", 11: "-"})
ws.append({1:"100-5", 2:"01.05.2024", 3: "АЕМ-1", 4: "12345-3456", 5: "Выпуск КД", 6: "Кошечкина"})
ws.append({1:"100-6", 2:"01.06.2024", 3: "АКГ-1", 4: "10800-0000", 5: "Выпуск КД", 6: "Кошечкина"})
ws.append({1:"100-7", 2: mydate.strftime("%d.%m.%Y"), 3: "АЕМ-1", 4: "12345-3456", 5: "Изготовление комплекта", 6: "Кошечкина"})

for i in range(3, 12):
    for j in range(1, 12):
        ws.cell(row=i, column=j).style = "other"

for i in range(1, 12):
    for j in range(1, 12):
        if i >= 3 and j == 5:
            ws.cell(row=i, column=5).alignment = Alignment(horizontal="left", vertical="center")
        else:
            ws.cell(row=i, column=j).alignment = Alignment(horizontal="center", vertical="center")

ws.auto_filter.ref = "A1:K11"
ws.auto_filter.add_filter_column(1, ["Кошечкина"])
ws.auto_filter.add_sort_condition("F3:F11")

wb.save("My_project.xlsx")
